#!/usr/bin/env python3
"""
Generate PowerPoint Presentation for SANZ GOKU Business Portfolio
Fully Styled to Match Persona 5 Aesthetic and Exact 8-Slide Requirements
Includes authentic background artwork, ransom note typography, polaroid photo,
and high-contrast Persona 5 UI cards.
"""

import os
import pptx
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw, ImageFont

# Initialize Presentation (16:9 Widescreen)
prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Directories
PROJECT_DIR = "/home/s4nzcyber/.gemini/antigravity/scratch/persona5-style-portfolio"
ASSETS_DIR = os.path.join(PROJECT_DIR, "assets")
MENUS_DIR = os.path.join(ASSETS_DIR, "menus")
TMP_DIR = "/tmp/p5_pptx_build"
os.makedirs(TMP_DIR, exist_ok=True)

FONT_PATH = "/tmp/ArchivoBlack.ttf"
PHOTO_PATH = os.path.join(ASSETS_DIR, "me.jpg")

# Persona 5 Theme Palette
COLOR_BG_DARK = RGBColor(11, 11, 14)        # Pure ink black
COLOR_BG_CARD = RGBColor(15, 15, 20)        # Dark ink card
COLOR_RED_PRIMARY = RGBColor(230, 0, 18)    # P5 Iconic Crimson Red
COLOR_RED_DEEP = RGBColor(163, 0, 12)       # P5 Deep Wine Red
COLOR_WHITE = RGBColor(246, 244, 239)       # P5 Off-white / Cream
COLOR_TEXT_MUTED = RGBColor(175, 175, 190)  # Light grey
COLOR_ACCENT_YELLOW = RGBColor(241, 224, 90)# P5 Gold / Yellow star
COLOR_BORDER = RGBColor(60, 60, 72)         # Card border

FONT_TITLE = "Arial Black"
FONT_BODY = "Arial"

# ==========================================
# 1. ASSET PREPARATION (PIL)
# ==========================================
print("Generating Persona 5 assets with PIL...")

# A. Generate Ransom Note Title Images
def render_ransom_text(text, font_size=46):
    f = ImageFont.truetype(FONT_PATH, font_size)
    styles = [
        {'bg': (11, 11, 13), 'fg': (246, 244, 239), 'rot': -4, 'border': (255, 255, 255)},
        {'bg': (246, 244, 239), 'fg': (11, 11, 13), 'rot': 4, 'border': (11, 11, 13)},
        {'bg': (230, 0, 18), 'fg': (246, 244, 239), 'rot': -3, 'border': (255, 255, 255)},
        {'bg': (11, 11, 13), 'fg': (230, 0, 18), 'rot': 3, 'border': (230, 0, 18)},
    ]
    boxes = []
    style_idx = 0
    for ch in text:
        if ch == ' ':
            boxes.append(('space', int(font_size * 0.4)))
            continue
        st = styles[style_idx % len(styles)]
        style_idx += 1
        bbox = f.getbbox(ch)
        cw = max(bbox[2] - bbox[0], int(font_size * 0.45)) + int(font_size * 0.28)
        ch_h = bbox[3] - bbox[1] + int(font_size * 0.32)
        
        pad = 20
        box_img = Image.new('RGBA', (cw + pad * 2, ch_h + pad * 2), (0,0,0,0))
        d = ImageDraw.Draw(box_img)
        # shadow
        d.rectangle([pad + 5, pad + 5, pad + cw + 5, pad + ch_h + 5], fill=(0,0,0,190))
        # box
        d.rectangle([pad, pad, pad + cw, pad + ch_h], fill=st['bg'], outline=st['border'], width=2)
        # letter centered
        tx = pad + (cw - (bbox[2] - bbox[0])) // 2 - bbox[0]
        ty = pad + (ch_h - (bbox[3] - bbox[1])) // 2 - bbox[1]
        d.text((tx, ty), ch, font=f, fill=st['fg'])
        
        rot_img = box_img.rotate(st['rot'], resample=Image.BICUBIC, expand=True)
        boxes.append(('box', rot_img))
    
    total_w = sum(item if typ == 'space' else item.width - int(font_size * 0.16) for typ, item in boxes) + 30
    max_h = max(item.height for typ, item in boxes if typ == 'box') + 20
    res = Image.new('RGBA', (total_w, max_h), (0,0,0,0))
    cur_x = 10
    for typ, item in boxes:
        if typ == 'space':
            cur_x += item
        else:
            y = (max_h - item.height) // 2
            res.paste(item, (cur_x, y), item)
            cur_x += item.width - int(font_size * 0.16)
    return res

slide_titles_def = [
    (1, "SANZ GOKU", 64),
    (2, "PROFIL SANZ GOKU", 44),
    (3, "PRODUK & LAYANAN", 44),
    (4, "SISTEM OPERASIONAL", 44),
    (5, "PEMASARAN & CS", 44),
    (6, "TEKNOLOGI PENDUKUNG", 44),
    (7, "ROADMAP PENGEMBANGAN", 44),
    (8, "KESIMPULAN BISNIS", 44),
]

title_images = {}
for s_num, t_text, fsize in slide_titles_def:
    t_img = render_ransom_text(t_text, fsize)
    t_path = os.path.join(TMP_DIR, f"title_{s_num}.png")
    t_img.save(t_path)
    title_images[s_num] = t_path

# B. Generate Slide Backgrounds (1920x1080)
w_px, h_px = 1920, 1080

# 1. Slide 1 (Cover): home.jpg with dark comic polygon on left
im_home = Image.open(os.path.join(MENUS_DIR, "home.jpg")).resize((w_px, h_px), Image.LANCZOS)
mask1 = Image.new('RGBA', (w_px, h_px), (0,0,0,0))
d1 = ImageDraw.Draw(mask1)
d1.polygon([(0, 0), (1240, 0), (940, h_px), (0, h_px)], fill=(11, 11, 14, 230))
d1.polygon([(1240, 0), (1260, 0), (960, h_px), (940, h_px)], fill=(230, 0, 18, 255))
bg1 = im_home.convert('RGBA')
bg1.alpha_composite(mask1)
bg1_path = os.path.join(TMP_DIR, "bg_slide1.jpg")
bg1.convert('RGB').save(bg1_path, quality=95)

# 2. Slide 2: about.jpg (Morgana) with dark panel on right
im_about = Image.open(os.path.join(MENUS_DIR, "about.jpg")).resize((w_px, h_px), Image.LANCZOS)
mask2 = Image.new('RGBA', (w_px, h_px), (0,0,0,0))
d2 = ImageDraw.Draw(mask2)
d2.polygon([(480, 0), (w_px, 0), (w_px, h_px), (380, h_px)], fill=(11, 11, 14, 235))
d2.polygon([(460, 0), (480, 0), (380, h_px), (360, h_px)], fill=(230, 0, 18, 255))
bg2 = im_about.convert('RGBA')
bg2.alpha_composite(mask2)
bg2_path = os.path.join(TMP_DIR, "bg_slide2.jpg")
bg2.convert('RGB').save(bg2_path, quality=95)

# 3. Slide 3, 4, 7: Stripes & Halftone Base
bg_stripes = Image.new('RGB', (w_px, h_px), (14, 14, 17))
d_s = ImageDraw.Draw(bg_stripes)
for i in range(-500, w_px + 1200, 140):
    d_s.polygon([(i, 0), (i + 65, 0), (i - 300, h_px), (i - 365, h_px)], fill=(230, 0, 18))
    d_s.polygon([(i + 65, 0), (i + 78, 0), (i - 287, h_px), (i - 300, h_px)], fill=(163, 0, 12))

mask_s = Image.new('RGBA', (w_px, h_px), (0,0,0,0))
ds_m = ImageDraw.Draw(mask_s)
ds_m.rectangle([0, 0, w_px, h_px], fill=(11, 11, 14, 160))
for x in range(0, w_px, 28):
    for y in range(0, h_px, 28):
        ds_m.ellipse([x, y, x+3, y+3], fill=(0,0,0,90))

bg_s_comp = bg_stripes.convert('RGBA')
bg_s_comp.alpha_composite(mask_s)
bg_stripes_path = os.path.join(TMP_DIR, "bg_stripes.jpg")
bg_s_comp.convert('RGB').save(bg_stripes_path, quality=95)

# 4. Slide 5: contact.jpg (Futaba)
im_contact = Image.open(os.path.join(MENUS_DIR, "contact.jpg")).resize((w_px, h_px), Image.LANCZOS)
mask5 = Image.new('RGBA', (w_px, h_px), (0,0,0,0))
d5 = ImageDraw.Draw(mask5)
d5.rectangle([0, 0, w_px, h_px], fill=(11, 11, 14, 175))
bg5 = im_contact.convert('RGBA')
bg5.alpha_composite(mask5)
bg5_path = os.path.join(TMP_DIR, "bg_slide5.jpg")
bg5.convert('RGB').save(bg5_path, quality=95)

# 5. Slide 6: skills.jpg (Makoto)
im_skills = Image.open(os.path.join(MENUS_DIR, "skills.jpg")).resize((w_px, h_px), Image.LANCZOS)
mask6 = Image.new('RGBA', (w_px, h_px), (0,0,0,0))
d6 = ImageDraw.Draw(mask6)
d6.rectangle([0, 0, w_px, h_px], fill=(11, 11, 14, 175))
bg6 = im_skills.convert('RGBA')
bg6.alpha_composite(mask6)
bg6_path = os.path.join(TMP_DIR, "bg_slide6.jpg")
bg6.convert('RGB').save(bg6_path, quality=95)

# C. Generate Polaroid Photo for Slide 2
font_pol_t = ImageFont.truetype(FONT_PATH, 20)
font_pol_s = ImageFont.truetype(FONT_PATH, 13)
me_img = Image.open(PHOTO_PATH).resize((370, 370), Image.LANCZOS)
pw, ph = 420, 540
polaroid = Image.new('RGBA', (pw, ph), (246, 244, 239, 255))
pdraw = ImageDraw.Draw(polaroid)
polaroid.paste(me_img, (25, 25))
pdraw.rectangle([24, 24, 395, 395], outline=(17, 17, 17), width=2)
pdraw.text((pw // 2, 412), 'M. SANDIKA AKBAR', font=font_pol_t, fill=(11, 11, 13), anchor='mt')
pdraw.text((pw // 2, 448), 'FOUNDER · SANZ GOKU', font=font_pol_s, fill=(230, 0, 18), anchor='mt')
pdraw.text((pw // 2, 474), 'snz.web.id · @Sanzstrnokosmumerbot', font=font_pol_s, fill=(80, 80, 90), anchor='mt')

# Red tape
tape = Image.new('RGBA', (170, 48), (0,0,0,0))
tdraw = ImageDraw.Draw(tape)
tdraw.polygon([(8, 0), (162, 6), (170, 42), (0, 36)], fill=(230, 0, 18, 235))
tdraw.rectangle([14, 8, 156, 34], outline=(255, 255, 255, 60), width=1)
tape_rot = tape.rotate(-7, resample=Image.BICUBIC, expand=True)
polaroid.paste(tape_rot, (115, -6), tape_rot)

final_pol = polaroid.rotate(-2.5, resample=Image.BICUBIC, expand=True)
shadow_pol = Image.new('RGBA', (final_pol.width + 30, final_pol.height + 30), (0,0,0,0))
sdraw = ImageDraw.Draw(shadow_pol)
sdraw.rectangle([15, 15, final_pol.width + 8, final_pol.height + 8], fill=(0,0,0,150))
shadow_pol.paste(final_pol, (5, 5), final_pol)
polaroid_path = os.path.join(TMP_DIR, "polaroid_me.png")
shadow_pol.save(polaroid_path)

print("Asset generation complete.")

# ==========================================
# 2. PPTX HELPER FUNCTIONS
# ==========================================
def apply_slide_bg(slide, bg_path):
    """Set full bleed slide background"""
    slide.shapes.add_picture(bg_path, 0, 0, Inches(13.333), Inches(7.5))

def add_header(slide, slide_num, category_str):
    """Add Persona 5 header with red badge, ransom title image, and slash line"""
    # 1. Red top strip
    top_strip = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
    top_strip.fill.solid()
    top_strip.fill.fore_color.rgb = COLOR_RED_PRIMARY
    top_strip.line.fill.background()

    # 2. Red Category Badge
    badge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.35), Inches(3.2), Inches(0.36))
    badge.fill.solid()
    badge.fill.fore_color.rgb = COLOR_RED_PRIMARY
    badge.line.fill.background()
    tf_b = badge.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = Inches(0.12)
    tf_b.margin_top = Inches(0.04)
    p_b = tf_b.paragraphs[0]
    p_b.text = f"{slide_num:02d} // {category_str}"
    p_b.font.name = FONT_TITLE
    p_b.font.size = Pt(11)
    p_b.font.bold = True
    p_b.font.color.rgb = COLOR_WHITE

    # 3. Ransom Note Title Image
    t_path = title_images[slide_num]
    # compute aspect ratio of image
    with Image.open(t_path) as im:
        iw, ih = im.size
    target_h = Inches(0.62)
    target_w = target_h * (iw / ih)
    slide.shapes.add_picture(t_path, Inches(0.78), Inches(0.76), target_w, target_h)

    # 4. Red slash accent line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.48), Inches(11.733), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_RED_PRIMARY
    line.line.fill.background()

def add_footer(slide, current_idx):
    """Add authentic HUD footer bar"""
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.35))
    tf_f = footer_box.text_frame
    tf_f.word_wrap = True
    p_f = tf_f.paragraphs[0]
    p_f.text = f"SANZ GOKU  ·  Satu Bot, Banyak Solusi  ·  snz.web.id                           SLIDE {current_idx:02d} / 08"
    p_f.font.name = FONT_BODY
    p_f.font.size = Pt(10)
    p_f.font.color.rgb = COLOR_TEXT_MUTED

# ==========================================
# SLIDE 1 — COVER
# ==========================================
slide1 = prs.slides.add_slide(blank_layout)
apply_slide_bg(slide1, bg1_path)

# Top Red Strip
s1_strip = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
s1_strip.fill.solid()
s1_strip.fill.fore_color.rgb = COLOR_RED_PRIMARY
s1_strip.line.fill.background()

# Top Tag
top_tag = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.7), Inches(3.6), Inches(0.42))
top_tag.fill.solid()
top_tag.fill.fore_color.rgb = COLOR_RED_PRIMARY
top_tag.line.fill.background()
tf_tag = top_tag.text_frame
p_tag = tf_tag.paragraphs[0]
p_tag.text = "PORTOFOLIO BISNIS DIGITAL"
p_tag.font.name = FONT_TITLE
p_tag.font.size = Pt(12)
p_tag.font.bold = True
p_tag.font.color.rgb = COLOR_WHITE

# Main Ransom Title: SANZ GOKU
t1_path = title_images[1]
with Image.open(t1_path) as im:
    iw, ih = im.size
target_h1 = Inches(1.15)
target_w1 = target_h1 * (iw / ih)
slide1.shapes.add_picture(t1_path, Inches(0.78), Inches(1.3), target_w1, target_h1)

# Tagline Badge
tagline_box = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(2.7), Inches(5.8), Inches(0.58))
tagline_box.fill.solid()
tagline_box.fill.fore_color.rgb = COLOR_BG_CARD
tagline_box.line.color.rgb = COLOR_ACCENT_YELLOW
tagline_box.line.width = Pt(1.5)
tf_tb = tagline_box.text_frame
tf_tb.margin_top = Inches(0.12)
p_tb = tf_tb.paragraphs[0]
p_tb.text = "« Satu Bot, Banyak Solusi »"
p_tb.font.name = FONT_TITLE
p_tb.font.size = Pt(16)
p_tb.font.bold = True
p_tb.font.color.rgb = COLOR_ACCENT_YELLOW

# Description Box
desc_box = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.5), Inches(7.6), Inches(1.6))
desc_box.fill.solid()
desc_box.fill.fore_color.rgb = COLOR_BG_CARD
desc_box.line.color.rgb = COLOR_BORDER
desc_box.line.width = Pt(1)
tf_dc = desc_box.text_frame
tf_dc.word_wrap = True
tf_dc.margin_left = Inches(0.25)
tf_dc.margin_top = Inches(0.2)
p_dc = tf_dc.paragraphs[0]
p_dc.text = "Penyedia produk & layanan digital, infrastruktur server Linux, panel Pterodactyl, nomor virtual OTP, serta solusi otomatisasi transaksi berbasis bot Telegram (@Sanzstrnokosmumerbot)."
p_dc.font.name = FONT_BODY
p_dc.font.size = Pt(13.5)
p_dc.font.color.rgb = COLOR_WHITE

# Presenter / Identity Box
id_box = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.3), Inches(7.6), Inches(1.45))
id_box.fill.solid()
id_box.fill.fore_color.rgb = COLOR_BG_CARD
id_box.line.color.rgb = COLOR_RED_PRIMARY
id_box.line.width = Pt(2)
tf_id = id_box.text_frame
tf_id.margin_left = Inches(0.25)
tf_id.margin_top = Inches(0.2)
p_id1 = tf_id.paragraphs[0]
p_id1.text = "DISUSUN OLEH:"
p_id1.font.name = FONT_TITLE
p_id1.font.size = Pt(11)
p_id1.font.color.rgb = COLOR_ACCENT_YELLOW
p_id2 = tf_id.add_paragraph()
p_id2.text = "Muhammad Sandika Akbar  (Sanz)"
p_id2.font.name = FONT_TITLE
p_id2.font.size = Pt(15)
p_id2.font.bold = True
p_id2.font.color.rgb = COLOR_WHITE
p_id3 = tf_id.add_paragraph()
p_id3.text = "Domain: snz.web.id  ·  Bot: @Sanzstrnokosmumerbot  ·  Admin: @Sanzxgoku"
p_id3.font.name = FONT_BODY
p_id3.font.size = Pt(11)
p_id3.font.color.rgb = COLOR_TEXT_MUTED

add_footer(slide1, 1)

# ==========================================
# SLIDE 2 — PROFIL BISNIS & OWNER
# ==========================================
slide2 = prs.slides.add_slide(blank_layout)
apply_slide_bg(slide2, bg2_path)
add_header(slide2, 2, "IDENTITAS & NILAI")

# Left Column: Authentic Polaroid Photo
with Image.open(polaroid_path) as im:
    pw_px, ph_px = im.size
pol_w = Inches(3.6)
pol_h = pol_w * (ph_px / pw_px)
slide2.shapes.add_picture(polaroid_path, Inches(0.7), Inches(1.7), pol_w, pol_h)

# Right Column: Executive Summary + 4 Pillars
# Executive Summary
card_top = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.5), Inches(1.7), Inches(8.033), Inches(1.35))
card_top.fill.solid()
card_top.fill.fore_color.rgb = COLOR_BG_CARD
card_top.line.color.rgb = COLOR_RED_PRIMARY
card_top.line.width = Pt(1.5)
tf_ct = card_top.text_frame
tf_ct.word_wrap = True
tf_ct.margin_left = Inches(0.25)
tf_ct.margin_right = Inches(0.25)
tf_ct.margin_top = Inches(0.15)
p_ct1 = tf_ct.paragraphs[0]
p_ct1.text = "★ RINGKASAN EKSEKUTIF"
p_ct1.font.name = FONT_TITLE
p_ct1.font.size = Pt(11)
p_ct1.font.bold = True
p_ct1.font.color.rgb = COLOR_ACCENT_YELLOW
p_ct2 = tf_ct.add_paragraph()
p_ct2.text = "SANZ GOKU adalah usaha digital aktif yang berfokus pada penyediaan produk & layanan digital, layanan server, serta solusi teknologi praktis. Usaha ini hadir untuk menjawab kebutuhan pelanggan akan infrastruktur hosting stabil, nomor virtual OTP cepat via bot Telegram, dan alur transaksi yang jelas serta transparan."
p_ct2.font.name = FONT_BODY
p_ct2.font.size = Pt(11)
p_ct2.font.color.rgb = COLOR_WHITE

# 4 Key Pillars (2x2 grid)
pillars = [
    ("BIDANG USAHA", [
        "Penjualan & konfigurasi VPS Linux.",
        "Instalasi panel Pterodactyl & hosting.",
        "Layanan nomor virtual OTP (RumahOTP).",
        "Script, subdomain, dan produk digital."
    ]),
    ("KEBUTUHAN PELANGGAN", [
        "Infrastruktur server fleksibel & stabil.",
        "Verifikasi akun instan tanpa nomor pribadi.",
        "Kebutuhan tools & script digital praktis.",
        "Pemesanan responsif dibantu admin."
    ]),
    ("TUJUAN PENDIRIAN", [
        "Solusi teknologi terjangkau & mudah diakses.",
        "Ekosistem mandiri berbasis efisiensi server.",
        "Otomatisasi pelayanan pelanggan bertahap.",
        "Penyedia layanan digital yang dapat diandalkan."
    ]),
    ("NILAI UTAMA BISNIS", [
        "Transparansi spesifikasi tanpa klaim palsu.",
        "Kecepatan penanganan pesanan via Telegram.",
        "Inovasi bot Telegram (@Sanzstrnokosmumerbot).",
        "Dukungan teknis responsif saat ada kendala."
    ]),
]

xp = [Inches(4.5), Inches(8.6)]
yp = [Inches(3.25), Inches(5.1)]

for idx, (p_title, p_items) in enumerate(pillars):
    col = idx % 2
    row = idx // 2
    w_box = Inches(3.933)
    h_box = Inches(1.72)
    
    box = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, xp[col], yp[row], w_box, h_box)
    box.fill.solid()
    box.fill.fore_color.rgb = COLOR_BG_CARD
    box.line.color.rgb = COLOR_BORDER
    box.line.width = Pt(1)
    
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.12)
    
    p0 = tf.paragraphs[0]
    p0.text = f"■  {p_title}"
    p0.font.name = FONT_TITLE
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_RED_PRIMARY
    
    for item in p_items:
        pi = tf.add_paragraph()
        pi.text = f"• {item}"
        pi.font.name = FONT_BODY
        pi.font.size = Pt(9.5)
        pi.font.color.rgb = COLOR_WHITE

add_footer(slide2, 2)

# ==========================================
# SLIDE 3 — PRODUK DAN LAYANAN
# ==========================================
slide3 = prs.slides.add_slide(blank_layout)
apply_slide_bg(slide3, bg_stripes_path)
add_header(slide3, 3, "KATALOG BISNIS")

# Status Warning Bar
status_bar = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.7), Inches(11.733), Inches(0.42))
status_bar.fill.solid()
status_bar.fill.fore_color.rgb = COLOR_BG_CARD
status_bar.line.color.rgb = COLOR_ACCENT_YELLOW
status_bar.line.width = Pt(1)
tf_sb = status_bar.text_frame
tf_sb.margin_left = Inches(0.25)
tf_sb.margin_top = Inches(0.08)
p_sb = tf_sb.paragraphs[0]
p_sb.text = "★ CATATAN STATUS: Layanan dikelompokkan berdasarkan kategori. Membedakan layanan aktif berjualan vs tahap rencana/pengembangan."
p_sb.font.name = FONT_BODY
p_sb.font.size = Pt(10)
p_sb.font.color.rgb = COLOR_ACCENT_YELLOW

services = [
    ("VPS & SERVER LINUX", "STATUS: AKTIF BERJUALAN", COLOR_RED_PRIMARY, [
        "Penjualan VPS Linux (berbagai pilihan spek).",
        "Setup & lisensi/manajemen panel Pterodactyl.",
        "Pengelolaan konfigurasi dasar server Linux.",
        "Hosting server sesuai ketersediaan stok."
    ]),
    ("NOKOS & LAYANAN OTP", "STATUS: AKTIF BERJUALAN", COLOR_RED_PRIMARY, [
        "Nomor virtual verifikasi akun aplikasi digital.",
        "Tersedia melalui Bot @Sanzstrnokosmumerbot.",
        "Integrasi penyedia terpercaya (RumahOTP).",
        "Aktivasi akun aman tanpa nomor pribadi."
    ]),
    ("PRODUK DIGITAL & LAINNYA", "STATUS: AKTIF / ON-DEMAND", COLOR_RED_PRIMARY, [
        "Penjualan script dan produk digital terpilih.",
        "Layanan penyediaan custom subdomain.",
        "Layanan boost media sosial & kebutuhan akun.",
        "Top up saldo dan layanan digital pelengkap."
    ]),
    ("BOT TELEGRAM (@Sanzstrnokosmumerbot)", "STATUS: PENGEMBANGAN FITUR", COLOR_ACCENT_YELLOW, [
        "Katalog interaktif otomatis via Telegram.",
        "Pemeriksaan status transaksi & saldo real-time.",
        "Rencana integrasi payment gateway QRIS.",
        "Mempermudah alur pemesanan secara mandiri."
    ]),
]

w_card = Inches(2.78)
gap = Inches(0.2)
y_card = Inches(2.28)
h_card = Inches(4.55)

for i, (title_s, status_s, badge_color, items_s) in enumerate(services):
    x_card = Inches(0.8) + i * (w_card + gap)
    card = slide3.shapes.add_shape(MSO_SHAPE.RECTANGLE, x_card, y_card, w_card, h_card)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_BG_CARD
    card.line.color.rgb = COLOR_RED_PRIMARY if badge_color == COLOR_RED_PRIMARY else COLOR_ACCENT_YELLOW
    card.line.width = Pt(1.5)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18)
    tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.18)
    
    p1 = tf.paragraphs[0]
    p1.text = title_s
    p1.font.name = FONT_TITLE
    p1.font.size = Pt(11.5)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    
    p2 = tf.add_paragraph()
    p2.text = f"[{status_s}]"
    p2.font.name = FONT_TITLE
    p2.font.size = Pt(9)
    p2.font.bold = True
    p2.font.color.rgb = badge_color
    
    p3 = tf.add_paragraph()
    p3.text = " "
    p3.font.size = Pt(4)
    
    for it in items_s:
        p_it = tf.add_paragraph()
        p_it.text = f"• {it}"
        p_it.font.name = FONT_BODY
        p_it.font.size = Pt(10.5)
        p_it.font.color.rgb = COLOR_WHITE
        p_it.space_after = Pt(6)

add_footer(slide3, 3)

# ==========================================
# SLIDE 4 — SISTEM OPERASIONAL BISNIS
# ==========================================
slide4 = prs.slides.add_slide(blank_layout)
apply_slide_bg(slide4, bg_stripes_path)
add_header(slide4, 4, "ALUR PELAYANAN")

box_op = slide4.shapes.add_textbox(Inches(0.75), Inches(1.62), Inches(11.8), Inches(0.5))
tf_op = box_op.text_frame
p_op = tf_op.paragraphs[0]
p_op.text = "Alur Pelayanan Pelanggan 6 Langkah — Menggabungkan penanganan admin aktif dengan persiapan otomasi bot (@Sanzstrnokosmumerbot)."
p_op.font.name = FONT_BODY
p_op.font.size = Pt(11.5)
p_op.font.color.rgb = COLOR_WHITE

steps = [
    ("01", "KATALOG & INFO", "Pelanggan melihat daftar layanan, deskripsi, dan ketersediaan stok melalui kanal resmi atau bot Telegram."),
    ("02", "PEMILIHAN LAYANAN", "Pelanggan memilih jenis layanan (VPS/OTP/Produk Digital) dan menyampaikan kebutuhan ke admin/bot."),
    ("03", "VERIFIKASI ADMIN", "Admin memeriksa rincian pesanan, kesiapan server/nomor, dan memberikan instruksi pembayaran."),
    ("04", "PEMBAYARAN", "Pelanggan melakukan pembayaran melalui metode yang tersedia dan diverifikasi secara transparan."),
    ("05", "PEMROSESAN SISTEM", "Pesanan diproses langsung melalui sistem internal server Linux atau penyedia layanan terkait."),
    ("06", "SERAH TERIMA & CS", "Hasil layanan diserahkan ke pelanggan (kredensial/kode OTP) disertai bantuan jika berkendala."),
]

for i, (num_s, title_step, desc_step) in enumerate(steps):
    col = i % 3
    row = i // 3
    x_s = Inches(0.8) + col * Inches(3.98)
    y_s = Inches(2.25) + row * Inches(2.15)
    w_s = Inches(3.78)
    h_s = Inches(1.95)
    
    card = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, x_s, y_s, w_s, h_s)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_BG_CARD
    card.line.color.rgb = COLOR_RED_PRIMARY
    card.line.width = Pt(1.5)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.15)
    
    p_num = tf.paragraphs[0]
    p_num.text = f"LANGKAH {num_s}"
    p_num.font.name = FONT_TITLE
    p_num.font.size = Pt(11)
    p_num.font.bold = True
    p_num.font.color.rgb = COLOR_RED_PRIMARY
    
    p_t = tf.add_paragraph()
    p_t.text = title_step
    p_t.font.name = FONT_TITLE
    p_t.font.size = Pt(12)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_WHITE
    
    p_d = tf.add_paragraph()
    p_d.text = desc_step
    p_d.font.name = FONT_BODY
    p_d.font.size = Pt(10)
    p_d.font.color.rgb = COLOR_TEXT_MUTED

note_box = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.58), Inches(11.733), Inches(0.38))
note_box.fill.solid()
note_box.fill.fore_color.rgb = COLOR_BG_CARD
note_box.line.color.rgb = COLOR_BORDER
tf_nb = note_box.text_frame
tf_nb.margin_left = Inches(0.2)
tf_nb.margin_top = Inches(0.06)
p_nb = tf_nb.paragraphs[0]
p_nb.text = "★ STATUS OPERASIONAL: Saat ini pesanan ditangani via admin. Integrasi Bot @Sanzstrnokosmumerbot terus disempurnakan untuk otomasi alur transaksi."
p_nb.font.name = FONT_BODY
p_nb.font.size = Pt(9.5)
p_nb.font.color.rgb = COLOR_ACCENT_YELLOW

add_footer(slide4, 4)

# ==========================================
# SLIDE 5 — PEMASARAN & PELAYANAN
# ==========================================
slide5 = prs.slides.add_slide(blank_layout)
apply_slide_bg(slide5, bg5_path)
add_header(slide5, 5, "PEMASARAN & PELAYANAN")

# Left Column: Pemasaran
box_mkt = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.75), Inches(5.733), Inches(4.0))
box_mkt.fill.solid()
box_mkt.fill.fore_color.rgb = COLOR_BG_CARD
box_mkt.line.color.rgb = COLOR_RED_PRIMARY
box_mkt.line.width = Pt(1.5)
tf_m = box_mkt.text_frame
tf_m.word_wrap = True
tf_m.margin_left = Inches(0.25)
tf_m.margin_right = Inches(0.25)
tf_m.margin_top = Inches(0.18)

pm_head = tf_m.paragraphs[0]
pm_head.text = "STRATEGI PEMASARAN DIGITAL"
pm_head.font.name = FONT_TITLE
pm_head.font.size = Pt(13)
pm_head.font.bold = True
pm_head.font.color.rgb = COLOR_RED_PRIMARY

mkt_points = [
    ("Pemanfaatan Kanal Digital", "Mengenalkan produk dan layanan melalui kanal media sosial dan komunitas online teknologi/gaming."),
    ("Katalog Produk yang Jelas", "Menyajikan spesifikasi layanan, batas penggunaan, serta panduan pemesanan tanpa klaim berlebihan."),
    ("Optimalisasi Telegram", "Menjadikan bot (@Sanzstrnokosmumerbot) dan chat admin (@Sanzxgoku) sebagai pusat interaksi harian yang cepat."),
    ("Harga Kompetitif & Realistis", "Menetapkan harga rasional sesuai spesifikasi server dan biaya penyedia tanpa perang harga merusak.")
]

for title_m, desc_m in mkt_points:
    p_t = tf_m.add_paragraph()
    p_t.text = f"■ {title_m}"
    p_t.font.name = FONT_TITLE
    p_t.font.size = Pt(11)
    p_t.font.color.rgb = COLOR_WHITE
    p_t.space_before = Pt(6)
    
    p_d = tf_m.add_paragraph()
    p_d.text = desc_m
    p_d.font.name = FONT_BODY
    p_d.font.size = Pt(9.5)
    p_d.font.color.rgb = COLOR_TEXT_MUTED

# Right Column: Pelayanan Pelanggan
box_cs = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.75), Inches(5.733), Inches(4.0))
box_cs.fill.solid()
box_cs.fill.fore_color.rgb = COLOR_BG_CARD
box_cs.line.color.rgb = COLOR_RED_PRIMARY
box_cs.line.width = Pt(1.5)
tf_c = box_cs.text_frame
tf_c.word_wrap = True
tf_c.margin_left = Inches(0.25)
tf_c.margin_right = Inches(0.25)
tf_c.margin_top = Inches(0.18)

pc_head = tf_c.paragraphs[0]
pc_head.text = "PELAYANAN & KEPERCAYAAN PELANGGAN"
pc_head.font.name = FONT_TITLE
pc_head.font.size = Pt(13)
pc_head.font.bold = True
pc_head.font.color.rgb = COLOR_RED_PRIMARY

cs_points = [
    ("Komunikasi Responsif", "Admin merespon pesan, pertanyaan spesifikasi, dan pesanan secara cepat serta lugas."),
    ("Bantuan & Troubleshooting", "Membantu pelanggan jika menemui kendala instalasi panel atau kendala kode OTP yang belum masuk."),
    ("Transparansi Stok & Server", "Memberitahukan ketersediaan kapasitas server apa adanya demi menjaga kepuasan pelanggan."),
    ("Membangun Loyalitas", "Pelayanan yang jujur dan tuntas menjadi fondasi repeat order tanpa perlu manipulasi testimoni.")
]

for title_c, desc_c in cs_points:
    p_t = tf_c.add_paragraph()
    p_t.text = f"■ {title_c}"
    p_t.font.name = FONT_TITLE
    p_t.font.size = Pt(11)
    p_t.font.color.rgb = COLOR_WHITE
    p_t.space_before = Pt(6)
    
    p_d = tf_c.add_paragraph()
    p_d.text = desc_c
    p_d.font.name = FONT_BODY
    p_d.font.size = Pt(9.5)
    p_d.font.color.rgb = COLOR_TEXT_MUTED

# Bottom Contact Bar
contact_box = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(5.95), Inches(11.733), Inches(0.85))
contact_box.fill.solid()
contact_box.fill.fore_color.rgb = COLOR_BG_CARD
contact_box.line.color.rgb = COLOR_ACCENT_YELLOW
contact_box.line.width = Pt(1.5)
tf_cb = contact_box.text_frame
tf_cb.word_wrap = True
tf_cb.margin_left = Inches(0.3)
tf_cb.margin_top = Inches(0.12)

p_cb1 = tf_cb.paragraphs[0]
p_cb1.text = "★ KANAL RESMI KOMUNIKASI & TRANSAKSI:"
p_cb1.font.name = FONT_TITLE
p_cb1.font.size = Pt(11)
p_cb1.font.color.rgb = COLOR_ACCENT_YELLOW

p_cb2 = tf_cb.add_paragraph()
p_cb2.text = "Bot Telegram: @Sanzstrnokosmumerbot  |  Admin: @Sanzxgoku  |  Instagram: @sanzshun  |  Domain: snz.web.id"
p_cb2.font.name = FONT_TITLE
p_cb2.font.size = Pt(12)
p_cb2.font.bold = True
p_cb2.font.color.rgb = COLOR_WHITE

add_footer(slide5, 5)

# ==========================================
# SLIDE 6 — TEKNOLOGI PENDUKUNG
# ==========================================
slide6 = prs.slides.add_slide(blank_layout)
apply_slide_bg(slide6, bg6_path)
add_header(slide6, 6, "STACK & INFRASTRUKTUR")

box_act = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.75), Inches(5.733), Inches(5.05))
box_act.fill.solid()
box_act.fill.fore_color.rgb = COLOR_BG_CARD
box_act.line.color.rgb = COLOR_RED_PRIMARY
box_act.line.width = Pt(1.5)
tf_a = box_act.text_frame
tf_a.word_wrap = True
tf_a.margin_left = Inches(0.25)
tf_a.margin_right = Inches(0.25)
tf_a.margin_top = Inches(0.18)

pa_h = tf_a.paragraphs[0]
pa_h.text = "TEKNOLOGI AKTIF DIGUNAKAN"
pa_h.font.name = FONT_TITLE
pa_h.font.size = Pt(13)
pa_h.font.bold = True
pa_h.font.color.rgb = COLOR_RED_PRIMARY

tech_active = [
    ("Linux VPS (Ubuntu / Debian)", "Infrastruktur dasar server untuk hosting layanan, menjalankan script, dan pengelolaan port jaringan."),
    ("JavaScript & Node.js", "Bahasa pemrograman dan runtime utama dalam membangun logika bot, skrip otomatisasi, serta komunikasi API."),
    ("Telegram Bot API", "Protokol penghubung untuk interaksi chat bot @Sanzstrnokosmumerbot, menu tombol, dan notifikasi."),
    ("Pterodactyl Panel", "Sistem panel web berbasis container Docker untuk mempermudah kontrol server game dan bot secara mandiri."),
    ("Cloudflare DNS & Proxy", "Manajemen domain snz.web.id, proteksi terhadap serangan DDoS, serta pengaturan routing subdomain.")
]

for t_name, t_desc in tech_active:
    p = tf_a.add_paragraph()
    p.text = f"■ {t_name}"
    p.font.name = FONT_TITLE
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_WHITE
    p.space_before = Pt(5)
    
    pd = tf_a.add_paragraph()
    pd.text = t_desc
    pd.font.name = FONT_BODY
    pd.font.size = Pt(9.5)
    pd.font.color.rgb = COLOR_TEXT_MUTED

box_pln = slide6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.8), Inches(1.75), Inches(5.733), Inches(5.05))
box_pln.fill.solid()
box_pln.fill.fore_color.rgb = COLOR_BG_CARD
box_pln.line.color.rgb = COLOR_ACCENT_YELLOW
box_pln.line.width = Pt(1.5)
tf_p = box_pln.text_frame
tf_p.word_wrap = True
tf_p.margin_left = Inches(0.25)
tf_p.margin_right = Inches(0.25)
tf_p.margin_top = Inches(0.18)

pp_h = tf_p.paragraphs[0]
pp_h.text = "INTEGRASI API & PROYEK PENGEMBANGAN"
pp_h.font.name = FONT_TITLE
pp_h.font.size = Pt(13)
pp_h.font.bold = True
pp_h.font.color.rgb = COLOR_ACCENT_YELLOW

tech_planned = [
    ("Integrasi API Pihak Ketiga (RumahOTP, FayuPedia)", "Rencana sinkronisasi stok dan eksekusi nomor OTP serta produk digital secara otomatis melalui REST API."),
    ("Pakasir (Gateway Pembayaran QRIS)", "Perencanaan sistem verifikasi pembayaran QRIS instan untuk bot Telegram dan web store."),
    ("Proyek SOTP (sxgotp)", "Pengembangan terpisah website layanan OTP modern menggunakan Next.js, React, Tailwind CSS, Vercel, dan Supabase."),
    ("Manfaat Utama Teknologi", "• Efisiensi Operasional: Mengurangi pekerjaan admin berulang.\n• Kecepatan Layanan: Transaksi diproses dalam hitungan detik.\n• Skalabilitas: Kapasitas sistem siap menangani lonjakan pesanan.")
]

for t_name, t_desc in tech_planned:
    p = tf_p.add_paragraph()
    p.text = f"◆ {t_name}"
    p.font.name = FONT_TITLE
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_WHITE
    p.space_before = Pt(5)
    
    pd = tf_p.add_paragraph()
    pd.text = t_desc
    pd.font.name = FONT_BODY
    pd.font.size = Pt(9.5)
    pd.font.color.rgb = COLOR_TEXT_MUTED

add_footer(slide6, 6)

# ==========================================
# SLIDE 7 — ROADMAP BISNIS
# ==========================================
slide7 = prs.slides.add_slide(blank_layout)
apply_slide_bg(slide7, bg_stripes_path)
add_header(slide7, 7, "ROADMAP BISNIS")

box_v = slide7.shapes.add_textbox(Inches(0.75), Inches(1.62), Inches(11.8), Inches(0.5))
tf_v = box_v.text_frame
p_v = tf_v.paragraphs[0]
p_v.text = "Rencana Strategis Masa Depan — Disajikan sebagai visi pengembangan jangka pendek dan menengah."
p_v.font.name = FONT_BODY
p_v.font.size = Pt(11.5)
p_v.font.color.rgb = COLOR_WHITE

plans = [
    ("01", "EKSPANSI LAYANAN DIGITAL", "Memperluas katalog produk digital, menambah varian spesifikasi server VPS, dan menyesuaikan layanan dengan tren kebutuhan pasar."),
    ("02", "PENINGKATAN MUTU LAYANAN", "Menjaga kestabilan uptime server, kecepatan delivery pesanan, serta evaluasi kepuasan pelanggan secara berkala."),
    ("03", "OTOMASI BOT (@Sanzstrnokosmumerbot)", "Menyelesaikan integrasi pemesanan mandiri, verifikasi QRIS otomatis, dan penyerahan kredensial/layanan langsung via bot."),
    ("04", "SISTEM PENCATATAN & PEMBUKUAN", "Menerapkan pencatatan transaksi terstruktur untuk memantau performa layanan, inventaris stok server, dan arus keuangan rapi."),
    ("05", "KEAMANAN & KEANDALAN SISTEM", "Memperkuat proteksi firewall server Linux, manajemen akses API yang aman, serta cadangan data (backup) rutin."),
    ("06", "EKOSISTEM BERKELANJUTAN", "Mengintegrasikan website (snz.web.id), bot Telegram, dan panel server menjadi satu ekosistem bisnis digital yang stabil.")
]

for i, (num_p, title_p, desc_p) in enumerate(plans):
    col = i % 2
    row = i // 2
    x_p = Inches(0.8) + col * Inches(6.0)
    y_p = Inches(2.25) + row * Inches(1.52)
    w_p = Inches(5.733)
    h_p = Inches(1.38)
    
    card = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, x_p, y_p, w_p, h_p)
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_BG_CARD
    card.line.color.rgb = COLOR_RED_PRIMARY
    card.line.width = Pt(1.5)
    
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.12)
    
    p_t = tf.paragraphs[0]
    p_t.text = f"★ {num_p}. {title_p}"
    p_t.font.name = FONT_TITLE
    p_t.font.size = Pt(11.5)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_ACCENT_YELLOW
    
    p_d = tf.add_paragraph()
    p_d.text = desc_p
    p_d.font.name = FONT_BODY
    p_d.font.size = Pt(9.5)
    p_d.font.color.rgb = COLOR_WHITE

add_footer(slide7, 7)

# ==========================================
# SLIDE 8 — PENUTUP & KESIMPULAN
# ==========================================
slide8 = prs.slides.add_slide(blank_layout)
apply_slide_bg(slide8, bg5_path)
add_header(slide8, 8, "PENUTUP")

# Big Executive Calling Card
card_close = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.75), Inches(11.733), Inches(2.65))
card_close.fill.solid()
card_close.fill.fore_color.rgb = COLOR_BG_CARD
card_close.line.color.rgb = COLOR_RED_PRIMARY
card_close.line.width = Pt(2)
tf_cl = card_close.text_frame
tf_cl.word_wrap = True
tf_cl.margin_left = Inches(0.35)
tf_cl.margin_right = Inches(0.35)
tf_cl.margin_top = Inches(0.25)

p_cl_tag = tf_cl.paragraphs[0]
p_cl_tag.text = "★ KESIMPULAN EKSEKUTIF"
p_cl_tag.font.name = FONT_TITLE
p_cl_tag.font.size = Pt(12)
p_cl_tag.font.bold = True
p_cl_tag.font.color.rgb = COLOR_ACCENT_YELLOW

p_cl_body = tf_cl.add_paragraph()
p_cl_body.text = "SANZ GOKU merupakan usaha digital yang berfokus pada produk dan layanan teknologi. Melalui pengembangan layanan server, produk digital, dan otomatisasi menggunakan bot Telegram (@Sanzstrnokosmumerbot), bisnis ini diarahkan untuk memberikan pelayanan yang praktis, transparan, dan terus berkembang sesuai kebutuhan pelanggan."
p_cl_body.font.name = FONT_BODY
p_cl_body.font.size = Pt(13)
p_cl_body.font.color.rgb = COLOR_WHITE
p_cl_body.space_before = Pt(8)

p_cl_motto = tf_cl.add_paragraph()
p_cl_motto.text = "“ SANZ GOKU — Satu Bot, Banyak Solusi. ”"
p_cl_motto.font.name = FONT_TITLE
p_cl_motto.font.size = Pt(16)
p_cl_motto.font.bold = True
p_cl_motto.font.color.rgb = COLOR_ACCENT_YELLOW
p_cl_motto.space_before = Pt(12)

# Bottom Row: Closing remark and Official Contacts
box_ty = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(4.65), Inches(4.5), Inches(2.15))
box_ty.fill.solid()
box_ty.fill.fore_color.rgb = COLOR_RED_PRIMARY
box_ty.line.fill.background()
tf_ty = box_ty.text_frame
tf_ty.word_wrap = True
tf_ty.margin_left = Inches(0.3)
tf_ty.margin_top = Inches(0.35)

p_ty1 = tf_ty.paragraphs[0]
p_ty1.text = "TERIMA KASIH"
p_ty1.font.name = FONT_TITLE
p_ty1.font.size = Pt(24)
p_ty1.font.bold = True
p_ty1.font.color.rgb = COLOR_WHITE

p_ty2 = tf_ty.add_paragraph()
p_ty2.text = "Presentasi Portofolio Bisnis Selesai."
p_ty2.font.name = FONT_BODY
p_ty2.font.size = Pt(12)
p_ty2.font.color.rgb = COLOR_WHITE
p_ty2.space_before = Pt(6)

p_ty3 = tf_ty.add_paragraph()
p_ty3.text = "Solusi Digital Cepat, Transparan, & Andal."
p_ty3.font.name = FONT_TITLE
p_ty3.font.size = Pt(11)
p_ty3.font.bold = True
p_ty3.font.color.rgb = COLOR_ACCENT_YELLOW
p_ty3.space_before = Pt(4)

box_ct = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.5), Inches(4.65), Inches(7.033), Inches(2.15))
box_ct.fill.solid()
box_ct.fill.fore_color.rgb = COLOR_BG_CARD
box_ct.line.color.rgb = COLOR_BORDER
box_ct.line.width = Pt(1)
tf_ct2 = box_ct.text_frame
tf_ct2.word_wrap = True
tf_ct2.margin_left = Inches(0.3)
tf_ct2.margin_top = Inches(0.18)

p_ct_t = tf_ct2.paragraphs[0]
p_ct_t.text = "INFORMASI KONTAK & IDENTITAS RESMI"
p_ct_t.font.name = FONT_TITLE
p_ct_t.font.size = Pt(12)
p_ct_t.font.bold = True
p_ct_t.font.color.rgb = COLOR_ACCENT_YELLOW

contacts = [
    ("Nama Pemilik", "Muhammad Sandika Akbar (Sanz)"),
    ("Bot Telegram", "@Sanzstrnokosmumerbot  (t.me/Sanzstrnokosmumerbot)"),
    ("Telegram Admin", "@Sanzxgoku  (t.me/Sanzxgoku)"),
    ("Instagram", "@sanzshun  (instagram.com/sanzshun)"),
    ("Domain Bisnis", "snz.web.id")
]

for label_c, val_c in contacts:
    p_row = tf_ct2.add_paragraph()
    p_row.text = f"• {label_c}: {val_c}"
    p_row.font.name = FONT_BODY
    p_row.font.size = Pt(10.5)
    p_row.font.color.rgb = COLOR_WHITE
    p_row.space_before = Pt(2)

add_footer(slide8, 8)

# ==========================================
# 3. SAVE PRESENTATION
# ==========================================
pptx_filename = "Portofolio_Bisnis_SANZ_GOKU.pptx"
output_path = os.path.join(PROJECT_DIR, pptx_filename)
prs.save(output_path)

# Also save copy to downloads directory
downloads_path = "/home/s4nzcyber/Unduhan/Portofolio_Bisnis_SANZ_GOKU.pptx"
try:
    prs.save(downloads_path)
    print(f"Copied to downloads: {downloads_path}")
except Exception as e:
    print(f"Note on downloads: {e}")

# Save secondary filename for compatibility
prs.save(os.path.join(PROJECT_DIR, "Portofolio_Bisnis_SANZ_GOKU_STORE.pptx"))

print(f"Presentation successfully created at: {output_path}")
print(f"Total slides: {len(prs.slides)}")
print("Persona 5 PPTX Generation Complete!")
